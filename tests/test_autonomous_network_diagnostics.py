from __future__ import annotations

import json
import threading
import time
from pathlib import Path

from openpyxl import load_workbook

from three_agent.security_monitoring.autonomous_diagnostics import (
    MAX_FAST_TCP_PORTS,
    AutonomousDiagnosticError,
    AutonomousNetworkDiagnosticAgent,
)
from three_agent.security_monitoring.collectors import CollectorResult
from three_agent.security_monitoring.contracts import AssetInventoryRecord, ObservationRecord
from three_agent.security_monitoring.policy import MonitoringPolicy, MonitoringPolicyEngine


class _ConcurrencyTracker:
    def __init__(self) -> None:
        self.active = 0
        self.max_active = 0
        self.lock = threading.Lock()

    def enter(self) -> None:
        with self.lock:
            self.active += 1
            self.max_active = max(self.max_active, self.active)

    def leave(self) -> None:
        with self.lock:
            self.active -= 1


class FakeIcmpCollector:
    def __init__(self, tracker: _ConcurrencyTracker | None = None, *, reachable: bool = True) -> None:
        self.tracker = tracker
        self.reachable = reachable

    def collect(self, *, asset, run_id, observed_at=None):
        if self.tracker:
            self.tracker.enter()
        try:
            time.sleep(0.03)
            observation = ObservationRecord(
                run_id=run_id,
                asset_id=asset.asset_id,
                collector="icmp_echo",
                observed_at=observed_at,
                metric="icmp_reachable",
                status="ok" if self.reachable else "unreachable",
                value=self.reachable,
                unit="bool",
            ).validate()
            return CollectorResult((observation,), None if self.reachable else "ICMP_UNREACHABLE")
        finally:
            if self.tracker:
                self.tracker.leave()


class FakeTcpCollector:
    def __init__(self, open_ports: set[int], tracker: _ConcurrencyTracker | None = None) -> None:
        self.open_ports = set(open_ports)
        self.tracker = tracker

    def collect(self, *, asset, port, run_id, observed_at=None):
        if self.tracker:
            self.tracker.enter()
        try:
            time.sleep(0.03)
            reachable = int(port) in self.open_ports
            observation = ObservationRecord(
                run_id=run_id,
                asset_id=asset.asset_id,
                collector="tcp_connect",
                observed_at=observed_at,
                metric=f"tcp_port_{int(port)}_reachable",
                status="ok" if reachable else "unreachable",
                value=reachable,
                unit="bool",
            ).validate()
            return CollectorResult((observation,), None if reachable else "TCP_CONNECT_FAILED")
        finally:
            if self.tracker:
                self.tracker.leave()


def _asset(*, ports=(80, 443, 554), caps=("icmp_echo", "tcp_connect")) -> AssetInventoryRecord:
    return AssetInventoryRecord(
        asset_id="cam-200",
        role="ip-camera",
        management_host="192.168.11.200",
        collector_capabilities=tuple(caps),
        allowed_tcp_ports=tuple(ports),
        data_class="confidential",
    ).validate()


def _engine(*, workers=4) -> MonitoringPolicyEngine:
    return MonitoringPolicyEngine(
        MonitoringPolicy(
            profile_id="diagnostic-test",
            allow_active_liveness=True,
            max_workers=workers,
            timeout_seconds=0.5,
        )
    )


def test_short_request_resolves_only_approved_inventory_target() -> None:
    agent = AutonomousNetworkDiagnosticAgent(
        _engine(),
        [_asset()],
        icmp_collector=FakeIcmpCollector(),
        tcp_collector=FakeTcpCollector({80}),
    )
    assert agent.resolve_asset("kiểm tra thiết bị 192.168.11.200").asset_id == "cam-200"
    assert agent.resolve_asset("check cam-200").management_host == "192.168.11.200"
    try:
        agent.resolve_asset("kiểm tra thiết bị 192.168.11.201")
    except AutonomousDiagnosticError as exc:
        assert str(exc) == "DIAGNOSTIC_TARGET_NOT_IN_APPROVED_INVENTORY"
    else:
        raise AssertionError("arbitrary target unexpectedly admitted")


def test_one_shot_diagnosis_runs_independent_probes_concurrently() -> None:
    tracker = _ConcurrencyTracker()
    agent = AutonomousNetworkDiagnosticAgent(
        _engine(workers=4),
        [_asset()],
        icmp_collector=FakeIcmpCollector(tracker),
        tcp_collector=FakeTcpCollector({80, 554}, tracker),
    )
    report = agent.diagnose("kiểm tra thiết bị 192.168.11.200")
    assert report.status == "online"
    assert report.confidence == "high"
    assert tracker.max_active >= 2
    assert {probe.probe_id for probe in report.probes} == {
        "icmp_echo",
        "tcp_connect:80",
        "tcp_connect:443",
        "tcp_connect:554",
    }
    assert any(finding.code == "HTTP_REACHABLE_HTTPS_NOT_REACHABLE" for finding in report.findings)


def test_report_generation_writes_json_and_real_xlsx(tmp_path: Path) -> None:
    agent = AutonomousNetworkDiagnosticAgent(
        _engine(),
        [_asset()],
        icmp_collector=FakeIcmpCollector(),
        tcp_collector=FakeTcpCollector({80}),
    )
    report = agent.diagnose(
        "kiểm tra thiết bị 192.168.11.200",
        report_dir=tmp_path,
    )
    assert report.report_json and Path(report.report_json).is_file()
    assert report.report_xlsx and Path(report.report_xlsx).is_file()
    payload = json.loads(Path(report.report_json).read_text(encoding="utf-8"))
    assert payload["status"] == "online"
    assert payload["authority"]["broad_port_scan"] is False
    workbook = load_workbook(report.report_xlsx, read_only=True)
    assert workbook.sheetnames == ["Summary", "Probes", "Findings", "Authority"]
    workbook.close()


def test_fast_path_never_expands_into_broad_port_scan() -> None:
    ports = tuple(range(10000, 10000 + MAX_FAST_TCP_PORTS + 4))
    agent = AutonomousNetworkDiagnosticAgent(
        _engine(),
        [_asset(ports=ports)],
        icmp_collector=FakeIcmpCollector(),
        tcp_collector=FakeTcpCollector(set()),
    )
    report = agent.diagnose("check 192.168.11.200")
    assert report.tested_tcp_ports == ports[:MAX_FAST_TCP_PORTS]
    assert report.skipped_tcp_ports == ports[MAX_FAST_TCP_PORTS:]
    assert any(finding.code == "FAST_PATH_PORT_BOUND_APPLIED" for finding in report.findings)


def test_no_approved_liveness_capability_fails_closed() -> None:
    asset = _asset(ports=(), caps=("local_net_read",))
    agent = AutonomousNetworkDiagnosticAgent(
        _engine(),
        [asset],
        icmp_collector=FakeIcmpCollector(),
        tcp_collector=FakeTcpCollector(set()),
    )
    try:
        agent.diagnose("check cam-200")
    except AutonomousDiagnosticError as exc:
        assert str(exc) == "DIAGNOSTIC_NO_APPROVED_LIVENESS_CAPABILITY"
    else:
        raise AssertionError("diagnostic unexpectedly ran without an approved liveness capability")
