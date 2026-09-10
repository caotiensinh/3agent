from __future__ import annotations

import hashlib
import ipaddress
import json
import re
import time
from concurrent.futures import ThreadPoolExecutor, as_completed
from dataclasses import asdict, dataclass
from datetime import datetime, timezone
from pathlib import Path
from typing import Iterable, Protocol

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font

from .collectors import CollectorResult, IcmpCollector, TcpConnectCollector
from .contracts import AssetInventoryRecord, ObservationRecord
from .policy import MonitoringPolicyEngine

AUTONOMOUS_DIAGNOSTIC_SCHEMA = "workspace-network-autonomous-diagnostic/v1"
MAX_DIAGNOSTIC_REQUEST_CHARS = 512
MAX_FAST_TCP_PORTS = 16


class AutonomousDiagnosticError(RuntimeError):
    """A one-shot diagnostic request is invalid or cannot be executed safely."""


class _IcmpProbe(Protocol):
    def collect(self, *, asset: AssetInventoryRecord, run_id: str, observed_at: str | None = None) -> CollectorResult: ...


class _TcpProbe(Protocol):
    def collect(self, *, asset: AssetInventoryRecord, port: int, run_id: str, observed_at: str | None = None) -> CollectorResult: ...


@dataclass(frozen=True)
class DiagnosticProbeResult:
    probe_id: str
    capability: str
    port: int | None
    status: str
    elapsed_ms: float
    failure_code: str | None
    observations: tuple[ObservationRecord, ...]

    def public_dict(self) -> dict[str, object]:
        return {
            "probe_id": self.probe_id,
            "capability": self.capability,
            "port": self.port,
            "status": self.status,
            "elapsed_ms": round(float(self.elapsed_ms), 3),
            "failure_code": self.failure_code,
            "observations": [asdict(item) for item in self.observations],
        }


@dataclass(frozen=True)
class DiagnosticFinding:
    severity: str
    code: str
    summary: str
    evidence: tuple[str, ...]

    def public_dict(self) -> dict[str, object]:
        return {
            "severity": self.severity,
            "code": self.code,
            "summary": self.summary,
            "evidence": list(self.evidence),
        }


@dataclass(frozen=True)
class AutonomousDiagnosticReport:
    request_sha256: str
    run_id: str
    generated_at: str
    asset_id: str
    target_host: str
    asset_role: str
    status: str
    confidence: str
    conclusion: str
    elapsed_ms: float
    probes: tuple[DiagnosticProbeResult, ...]
    findings: tuple[DiagnosticFinding, ...]
    tested_tcp_ports: tuple[int, ...]
    skipped_tcp_ports: tuple[int, ...]
    report_json: str | None = None
    report_xlsx: str | None = None
    schema_version: str = AUTONOMOUS_DIAGNOSTIC_SCHEMA

    def public_dict(self) -> dict[str, object]:
        return {
            "schema_version": self.schema_version,
            "request_sha256": self.request_sha256,
            "run_id": self.run_id,
            "generated_at": self.generated_at,
            "asset": {
                "asset_id": self.asset_id,
                "management_host": self.target_host,
                "role": self.asset_role,
            },
            "status": self.status,
            "confidence": self.confidence,
            "conclusion": self.conclusion,
            "elapsed_ms": round(float(self.elapsed_ms), 3),
            "tested_tcp_ports": list(self.tested_tcp_ports),
            "skipped_tcp_ports": list(self.skipped_tcp_ports),
            "probes": [item.public_dict() for item in self.probes],
            "findings": [item.public_dict() for item in self.findings],
            "reports": {
                "json": self.report_json,
                "xlsx": self.report_xlsx,
            },
            "authority": {
                "network_scope": "approved_inventory_only",
                "read_only": True,
                "arbitrary_target_input": False,
                "broad_port_scan": False,
                "credential_use": False,
                "configuration_write": False,
                "remediation_execution": False,
            },
        }


class AutonomousNetworkDiagnosticAgent:
    """Fast inventory-bound one-shot diagnostics for terse operator requests.

    The request selects an already-approved asset. It never expands authority: probes
    are limited to the asset's reviewed collector capabilities and explicit TCP port
    allowlist. Independent liveness probes execute concurrently within the monitoring
    policy worker bound. Conclusions are deterministic and evidence-derived.
    """

    def __init__(
        self,
        policy_engine: MonitoringPolicyEngine,
        assets: Iterable[AssetInventoryRecord],
        *,
        icmp_collector: _IcmpProbe | None = None,
        tcp_collector: _TcpProbe | None = None,
    ) -> None:
        self.policy_engine = policy_engine
        validated = tuple(asset.validate() for asset in assets)
        self.assets = tuple(asset for asset in validated if asset.enabled)
        self.icmp = icmp_collector or IcmpCollector(policy_engine)
        self.tcp = tcp_collector or TcpConnectCollector(policy_engine)

    @staticmethod
    def _request_sha256(request: str) -> str:
        return "sha256:" + hashlib.sha256(request.encode("utf-8")).hexdigest()

    def resolve_asset(self, request: str) -> AssetInventoryRecord:
        raw = str(request or "").strip()
        if not raw or len(raw) > MAX_DIAGNOSTIC_REQUEST_CHARS:
            raise AutonomousDiagnosticError("DIAGNOSTIC_REQUEST_INVALID")

        matches: list[AssetInventoryRecord] = []
        lowered = raw.casefold()
        for asset in self.assets:
            if re.search(rf"(?<![A-Za-z0-9_.:-]){re.escape(asset.asset_id.casefold())}(?![A-Za-z0-9_.:-])", lowered):
                matches.append(asset)
                continue
            host = asset.management_host
            if _request_mentions_host(raw, host):
                matches.append(asset)

        unique = {asset.asset_id: asset for asset in matches}
        if len(unique) == 1:
            return next(iter(unique.values()))
        if len(unique) > 1:
            raise AutonomousDiagnosticError("DIAGNOSTIC_TARGET_AMBIGUOUS")
        raise AutonomousDiagnosticError("DIAGNOSTIC_TARGET_NOT_IN_APPROVED_INVENTORY")

    def diagnose(
        self,
        request: str,
        *,
        report_dir: Path | str | None = None,
    ) -> AutonomousDiagnosticReport:
        asset = self.resolve_asset(request)
        started = time.perf_counter()
        generated_at = datetime.now(timezone.utc).isoformat()
        run_id = "diag-" + hashlib.sha256(
            f"{asset.asset_id}|{generated_at}|{request}".encode("utf-8")
        ).hexdigest()[:20]

        allowed_ports = tuple(sorted(set(int(port) for port in asset.allowed_tcp_ports)))
        tested_ports = allowed_ports[:MAX_FAST_TCP_PORTS]
        skipped_ports = allowed_ports[MAX_FAST_TCP_PORTS:]

        jobs: list[tuple[str, int | None]] = []
        if "icmp_echo" in asset.collector_capabilities:
            jobs.append(("icmp_echo", None))
        if "tcp_connect" in asset.collector_capabilities:
            jobs.extend(("tcp_connect", port) for port in tested_ports)
        if not jobs:
            raise AutonomousDiagnosticError("DIAGNOSTIC_NO_APPROVED_LIVENESS_CAPABILITY")

        probe_results: list[DiagnosticProbeResult] = []
        workers = min(self.policy_engine.policy.max_workers, len(jobs))
        with ThreadPoolExecutor(max_workers=max(1, workers), thread_name_prefix="workspace-diag") as pool:
            futures = {
                pool.submit(self._run_probe, asset, run_id, generated_at, capability, port): (capability, port)
                for capability, port in jobs
            }
            for future in as_completed(futures):
                probe_results.append(future.result())

        probe_results.sort(key=lambda item: (0 if item.capability == "icmp_echo" else 1, item.port or 0))
        status, confidence, conclusion, findings = _diagnose(asset, tuple(probe_results), skipped_ports)
        elapsed_ms = (time.perf_counter() - started) * 1000.0
        report = AutonomousDiagnosticReport(
            request_sha256=self._request_sha256(str(request).strip()),
            run_id=run_id,
            generated_at=generated_at,
            asset_id=asset.asset_id,
            target_host=asset.management_host,
            asset_role=asset.role,
            status=status,
            confidence=confidence,
            conclusion=conclusion,
            elapsed_ms=elapsed_ms,
            probes=tuple(probe_results),
            findings=findings,
            tested_tcp_ports=tested_ports,
            skipped_tcp_ports=skipped_ports,
        )
        if report_dir is None:
            return report
        return self._write_reports(report, Path(report_dir))

    def _run_probe(
        self,
        asset: AssetInventoryRecord,
        run_id: str,
        observed_at: str,
        capability: str,
        port: int | None,
    ) -> DiagnosticProbeResult:
        started = time.perf_counter()
        if capability == "icmp_echo":
            result = self.icmp.collect(asset=asset, run_id=run_id, observed_at=observed_at)
        elif capability == "tcp_connect" and port is not None:
            result = self.tcp.collect(asset=asset, port=port, run_id=run_id, observed_at=observed_at)
        else:
            raise AutonomousDiagnosticError("DIAGNOSTIC_PROBE_NOT_IMPLEMENTED")
        elapsed_ms = (time.perf_counter() - started) * 1000.0
        status = "ok" if result.failure_code is None else "failed"
        probe_id = capability if port is None else f"{capability}:{port}"
        return DiagnosticProbeResult(
            probe_id=probe_id,
            capability=capability,
            port=port,
            status=status,
            elapsed_ms=elapsed_ms,
            failure_code=result.failure_code,
            observations=result.observations,
        )

    def _write_reports(
        self,
        report: AutonomousDiagnosticReport,
        report_dir: Path,
    ) -> AutonomousDiagnosticReport:
        report_dir.mkdir(parents=True, exist_ok=True)
        safe_asset = re.sub(r"[^A-Za-z0-9._-]+", "_", report.asset_id)[:80]
        stem = f"network_diagnostic_{safe_asset}_{report.run_id[-8:]}"
        json_path = report_dir / f"{stem}.json"
        xlsx_path = report_dir / f"{stem}.xlsx"

        payload = report.public_dict()
        payload["reports"] = {"json": str(json_path), "xlsx": str(xlsx_path)}
        json_path.write_text(
            json.dumps(payload, ensure_ascii=False, sort_keys=True, indent=2) + "\n",
            encoding="utf-8",
        )
        _write_xlsx(report, xlsx_path)
        return AutonomousDiagnosticReport(
            **{
                **asdict(report),
                "report_json": str(json_path),
                "report_xlsx": str(xlsx_path),
                "probes": report.probes,
                "findings": report.findings,
                "tested_tcp_ports": report.tested_tcp_ports,
                "skipped_tcp_ports": report.skipped_tcp_ports,
            }
        )


def _request_mentions_host(request: str, host: str) -> bool:
    try:
        normalized = str(ipaddress.ip_address(host))
    except ValueError:
        return bool(re.search(rf"(?<![A-Za-z0-9_.-]){re.escape(host)}(?![A-Za-z0-9_.-])", request, re.IGNORECASE))
    candidates = re.findall(r"(?<![0-9.])(?:\d{1,3}\.){3}\d{1,3}(?![0-9.])", request)
    for candidate in candidates:
        try:
            if str(ipaddress.ip_address(candidate)) == normalized:
                return True
        except ValueError:
            continue
    return False


def _bool_metric(probes: tuple[DiagnosticProbeResult, ...], metric: str) -> bool | None:
    for probe in probes:
        for observation in probe.observations:
            if observation.metric == metric and isinstance(observation.value, bool):
                return observation.value
    return None


def _open_ports(probes: tuple[DiagnosticProbeResult, ...]) -> tuple[int, ...]:
    ports: list[int] = []
    for probe in probes:
        if probe.capability != "tcp_connect" or probe.port is None:
            continue
        if _bool_metric((probe,), f"tcp_port_{probe.port}_reachable") is True:
            ports.append(probe.port)
    return tuple(sorted(ports))


def _diagnose(
    asset: AssetInventoryRecord,
    probes: tuple[DiagnosticProbeResult, ...],
    skipped_ports: tuple[int, ...],
) -> tuple[str, str, str, tuple[DiagnosticFinding, ...]]:
    icmp_reachable = _bool_metric(probes, "icmp_reachable")
    open_ports = _open_ports(probes)
    findings: list[DiagnosticFinding] = []

    positive_count = (1 if icmp_reachable is True else 0) + len(open_ports)
    if positive_count:
        status = "online"
        confidence = "high" if positive_count >= 2 else "medium"
        conclusion = (
            f"Asset {asset.asset_id} is reachable. "
            f"Observed open approved TCP ports: {', '.join(map(str, open_ports)) or 'none'}; "
            f"ICMP reachable: {icmp_reachable is True}."
        )
    else:
        conclusive = [
            obs
            for probe in probes
            for obs in probe.observations
            if obs.metric == "icmp_reachable" or obs.metric.startswith("tcp_port_")
        ]
        has_unknown = any(obs.status in {"timeout", "unsupported", "error"} for obs in conclusive)
        if conclusive and not has_unknown and all(obs.value is False for obs in conclusive):
            status = "offline"
            confidence = "medium"
            conclusion = "No approved liveness probe succeeded; the asset is currently unreachable from this diagnostic node."
        else:
            status = "ambiguous"
            confidence = "low"
            conclusion = "Available read-only probes are insufficient to prove the asset online or offline."

    if 23 in open_ports:
        findings.append(
            DiagnosticFinding(
                "high",
                "TELNET_SERVICE_REACHABLE",
                "Telnet is reachable on an explicitly approved port; prefer an encrypted management protocol where supported.",
                ("tcp_connect:23",),
            )
        )
    if 80 in open_ports and 443 in asset.allowed_tcp_ports and 443 not in open_ports:
        findings.append(
            DiagnosticFinding(
                "medium",
                "HTTP_REACHABLE_HTTPS_NOT_REACHABLE",
                "HTTP is reachable while the approved HTTPS port did not accept a connection. This is transport evidence only; protocol configuration was not changed or authenticated.",
                ("tcp_connect:80", "tcp_connect:443"),
            )
        )
    if skipped_ports:
        findings.append(
            DiagnosticFinding(
                "info",
                "FAST_PATH_PORT_BOUND_APPLIED",
                f"Fast diagnostic mode tested the first {MAX_FAST_TCP_PORTS} approved TCP ports and left {len(skipped_ports)} additional approved ports for a deliberate deep diagnostic.",
                tuple(f"approved-port:{port}" for port in skipped_ports[:8]),
            )
        )
    if status == "offline":
        findings.append(
            DiagnosticFinding(
                "medium",
                "LIVENESS_FAILURE_REQUIRES_UPSTREAM_CORRELATION",
                "An unreachable result does not by itself distinguish device power loss, cabling/PoE failure, VLAN/routing failure, ACL policy, or host failure. Correlate upstream switch/gateway/power evidence before root-cause closure.",
                tuple(probe.probe_id for probe in probes),
            )
        )
    return status, confidence, conclusion, tuple(findings)


def _write_xlsx(report: AutonomousDiagnosticReport, path: Path) -> None:
    workbook = Workbook()
    summary = workbook.active
    summary.title = "Summary"
    summary.append(["WorkSpace Autonomous Network Diagnostic"])
    summary["A1"].font = Font(bold=True, size=16)
    summary.append([])
    rows = [
        ("Asset ID", report.asset_id),
        ("Target", report.target_host),
        ("Role", report.asset_role),
        ("Status", report.status.upper()),
        ("Confidence", report.confidence),
        ("Conclusion", report.conclusion),
        ("Generated At", report.generated_at),
        ("Elapsed (ms)", round(report.elapsed_ms, 3)),
        ("Tested TCP Ports", ", ".join(map(str, report.tested_tcp_ports)) or "None"),
        ("Skipped TCP Ports", ", ".join(map(str, report.skipped_tcp_ports)) or "None"),
    ]
    for key, value in rows:
        summary.append([key, value])
    for cell in summary["A"]:
        cell.font = Font(bold=True)
    summary.column_dimensions["A"].width = 24
    summary.column_dimensions["B"].width = 90
    summary["B7"].alignment = Alignment(wrap_text=True, vertical="top")

    probes_sheet = workbook.create_sheet("Probes")
    probes_sheet.append(["Probe", "Capability", "Port", "Status", "Elapsed ms", "Failure Code", "Metric", "Observation Status", "Value", "Unit"])
    for cell in probes_sheet[1]:
        cell.font = Font(bold=True)
    for probe in report.probes:
        if not probe.observations:
            probes_sheet.append([probe.probe_id, probe.capability, probe.port, probe.status, round(probe.elapsed_ms, 3), probe.failure_code, "", "", "", ""])
            continue
        for observation in probe.observations:
            probes_sheet.append([
                probe.probe_id,
                probe.capability,
                probe.port,
                probe.status,
                round(probe.elapsed_ms, 3),
                probe.failure_code,
                observation.metric,
                observation.status,
                observation.value,
                observation.unit,
            ])
    probes_sheet.freeze_panes = "A2"

    findings_sheet = workbook.create_sheet("Findings")
    findings_sheet.append(["Severity", "Code", "Summary", "Evidence"])
    for cell in findings_sheet[1]:
        cell.font = Font(bold=True)
    for finding in report.findings:
        findings_sheet.append([finding.severity, finding.code, finding.summary, ", ".join(finding.evidence)])
    findings_sheet.freeze_panes = "A2"
    findings_sheet.column_dimensions["C"].width = 100
    findings_sheet.column_dimensions["D"].width = 60

    evidence_sheet = workbook.create_sheet("Authority")
    evidence_sheet.append(["Control", "Value"])
    for cell in evidence_sheet[1]:
        cell.font = Font(bold=True)
    for key, value in report.public_dict()["authority"].items():
        evidence_sheet.append([key, value])

    workbook.save(path)
